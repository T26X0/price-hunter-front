from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path("outputs/Задание_1_TQM_SWOT.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="808080", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_col_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_text(paragraph, text, bold=False, size=11, first_line=0):
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.first_line_indent = Cm(first_line)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_numbered_item(document, number, text, before=2, after=2):
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.columns[0].width = Cm(0.72)
    table.columns[1].width = Cm(15.9)
    for idx, cell in enumerate(table.row_cells(0)):
        set_cell_margins(cell, top=0, start=0 if idx == 0 else 110, bottom=0, end=0)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    p_num = table.cell(0, 0).paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_num.paragraph_format.space_before = Pt(before)
    add_text(p_num, f"{number}.", size=11)
    p = table.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    add_text(p, text, size=11)
    return table


def add_comparison_table(document):
    headers = ["Элемент", "Традиционная\nмодель", "TQM", "Пример для\n«ПромТехМонтаж»"]
    elements = [
        "Организационная\nструктура",
        "Стиль руководства",
        "Центр внимания",
        "Мотивация",
        "Рынки",
        "Ресурсы",
        "Рабочая сила",
        "Ожидания сотрудников",
        "Организация работы",
    ]
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [5.0, 4.0, 1.05, 6.1]
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_col_width(cell, widths[idx])
        set_cell_margins(cell)
        set_cell_border(cell)
        set_cell_shading(cell, "F2F2F2")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, header, size=10)

    for element in elements:
        cells = table.add_row().cells
        for idx, cell in enumerate(cells):
            set_col_width(cell, widths[idx])
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cells[0].paragraphs[0]
        add_text(p, element, size=10)
        for idx in (1, 2, 3):
            add_text(cells[idx].paragraphs[0], "", size=10)
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    title_style_pr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_style_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_style_pr.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.line_spacing = 1.0
    run = title.add_run("Задание 1: Комплексный SWOT-анализ и разработка стратегии качества")
    run.bold = False
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0, 0, 0)

    intro = doc.add_paragraph()
    add_text(
        intro,
        "Компания «ПромТехМонтаж» (производство промышленного оборудования, 320 сотрудников) планирует внедрение TQM. "
        "Организация работает 12 лет, имеет линейно-функциональную структуру, авторитарный стиль руководства. "
        "За последние 3 года 18% продукции возвращено клиентами из-за несоответствия требованиям. При этом 60% затрат "
        "на качество приходится на исправление брака. Текучесть кадров — 28% в год.",
        size=11,
    )
    intro.paragraph_format.space_after = Pt(12)
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    task_heading = doc.add_paragraph()
    task_heading.paragraph_format.space_after = Pt(5)
    add_text(task_heading, "Задание:", bold=True, size=11)

    add_numbered_item(
        doc,
        1,
        "Проведите SWOT-анализ организации, заполнив таблицу с не менее чем 5 сильных сторон, 5 слабых, "
        "5 возможностей и 5 угроз. Каждый пункт обоснуйте конкретными фактами и логическими выводами.",
    )
    add_numbered_item(
        doc,
        2,
        "Составьте таблицу сравнения традиционной модели менеджмента и TQM по всем 9 элементам, приведённым "
        "в лекции, с конкретными примерами для данной организации:",
        after=4,
    )
    add_comparison_table(doc)

    p_after_table = doc.add_paragraph()
    p_after_table.paragraph_format.space_after = Pt(1)
    p_after_table.paragraph_format.space_before = Pt(5)

    add_numbered_item(
        doc,
        3,
        "Определите, какие из 5 «смертельных болезней» Э. Деминга присутствуют в данной организации. Для каждой "
        "обнаруженной болезни приведите доказательства из описанной ситуации и предложите конкретный способ её устранения.",
        before=0,
    )
    add_numbered_item(
        doc,
        4,
        "Постройте «цепную реакцию Деминга» для данной организации, начиная с текущего состояния качества и показывая, "
        "какие последствия возникают в настоящее время и что изменится при внедрении TQM.",
        before=1,
        after=0,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
